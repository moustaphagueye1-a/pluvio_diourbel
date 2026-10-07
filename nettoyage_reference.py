"""
Étape 2 - Sources de référence et calendrier.

  1. Tableau 2012-2025 : cumul annuel (mm) et jours de pluie par poste
  2. Normales 1991-2020 : cumul de référence par jour (1er mai -> 31 octobre)
  3. Calendrier des semaines de bulletin, avec le jour de début et de fin

À lancer APRÈS nettoyage.py (il lit la table observations_hebdo).

Utilisation :
    python nettoyage_reference.py chemin/vers/le_fichier.xlsx
"""
import db
import sys
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from nettoyage import REFERENTIEL, normaliser, vers_nombre

ANNEE_COURANTE = 2026
JOURS_FR = ["lundi", "mardi", "mercredi", "jeudi", "vendredi", "samedi",
            "dimanche"]
TOL = 0.06  # tolérance d'arrondi (les normales du bulletin sont arrondies)
alertes = []


def alerte(source, detail):
    alertes.append(dict(source=source, detail=detail))


# --------------------------------------------------------------------------
# 1. Tableau 2012-2025
# --------------------------------------------------------------------------
def lire_annuel(ws):
    """Renvoie (données par poste et année, moyennes écrites dans le fichier)."""
    # Ligne 6 : les années sont dans des cellules fusionnées (colonnes B, D,
    # F...). La colonne mm est celle de l'année, la colonne jours est la suivante.
    annees = {c: ws.cell(6, c).value for c in range(2, ws.max_column + 1)
              if isinstance(ws.cell(6, c).value, int)}
    lignes, moyennes, dept = [], {}, None
    for r in range(8, ws.max_row + 1):
        nom = ws.cell(r, 1).value
        if nom is None:
            continue
        cle = normaliser(nom)
        if cle in REFERENTIEL:
            poste, dept = REFERENTIEL[cle]
            for col, an in annees.items():
                mm, st = vers_nombre(ws.cell(r, col).value)
                nj, _ = vers_nombre(ws.cell(r, col + 1).value)
                lignes.append(dict(
                    poste=poste, departement=dept, annee=an, mm=mm,
                    jours_pluie=nj,
                    statut="mesure" if mm is not None else "manquant"))
                if mm is None:
                    alerte("annuel", f"{poste} {an} : pas de donnée")
        else:
            # Ligne de moyenne (MOY/DEPT, ligne au nom vide, MOY/REGION)
            etiquette = "Région" if "region" in cle else dept
            for col, an in annees.items():
                mm, _ = vers_nombre(ws.cell(r, col).value)
                nj, _ = vers_nombre(ws.cell(r, col + 1).value)
                moyennes[(etiquette, an)] = (mm, nj)
    return pd.DataFrame(lignes), moyennes


def tests_annuel(a, moyennes):
    res = []
    n = a.groupby("annee")["poste"].nunique()
    res.append(("19 postes listés par année", (n == 19).all()))
    # Recoupement : mes moyennes recalculées = moyennes écrites dans le fichier
    ecarts = []
    for (etiq, an), (mm_f, nj_f) in moyennes.items():
        sel = a[(a.annee == an) & a.mm.notna()]
        if etiq != "Région":
            sel = sel[sel.departement == etiq]
        if mm_f is None or sel.empty:
            continue
        ecarts.append(abs(sel.mm.mean() - mm_f))
        if nj_f is not None:
            ecarts.append(abs(sel.jours_pluie.mean() - nj_f) / 10)
    res.append((f"moyennes recalculées = moyennes du fichier "
                f"({len(ecarts)} comparaisons, écart max "
                f"{max(ecarts):.3f})", max(ecarts) < 0.01))
    res.append(("aucune pluie négative ni jours > 366",
                bool((a.mm.dropna() >= 0).all()
                     and (a.jours_pluie.dropna() <= 366).all())))
    return res


# --------------------------------------------------------------------------
# 2. Normales 1991-2020
# --------------------------------------------------------------------------
def lire_normales(ws):
    """La colonne « Date » est une vraie date Excel (année sans importance) :
    on garde le jour et le mois, et on les rattache à 2026."""
    lignes = []
    for r in ws.iter_rows(min_row=5, values_only=True):
        if not isinstance(r[0], datetime):
            continue
        d = date(ANNEE_COURANTE, r[0].month, r[0].day)
        for dept, v in zip(["Bambey", "Diourbel", "Mbacké"], r[1:4]):
            lignes.append(dict(departement=dept, date=d, mois=d.month,
                               jour=d.day, jour_semaine=JOURS_FR[d.weekday()],
                               cumul_normale_mm=float(v)))
    return pd.DataFrame(lignes)


def tests_normales(nm):
    res = []
    dates = sorted(nm["date"].unique())
    jours = [(b - a).days for a, b in zip(dates, dates[1:])]
    res.append((f"{len(dates)} dates consécutives "
                f"({dates[0]} -> {dates[-1]})", set(jours) == {1}))
    croissant = all(g.sort_values("date").cumul_normale_mm.diff().dropna()
                    .ge(-1e-9).all() for _, g in nm.groupby("departement"))
    res.append(("normale cumulée jamais décroissante", croissant))
    return res


# --------------------------------------------------------------------------
# 3. Calendrier des semaines + lien avec les normales
# --------------------------------------------------------------------------
def construire_calendrier(h):
    s = (h[["semaine_debut", "semaine_fin"]].drop_duplicates()
         .sort_values("semaine_debut").reset_index(drop=True))
    s["semaine_debut"] = pd.to_datetime(s.semaine_debut)
    s["semaine_fin"] = pd.to_datetime(s.semaine_fin)
    s.insert(0, "num_semaine", range(1, len(s) + 1))
    s["jour_debut"] = s.semaine_debut.dt.weekday.map(lambda i: JOURS_FR[i])
    s["jour_fin"] = s.semaine_fin.dt.weekday.map(lambda i: JOURS_FR[i])
    s["nb_jours"] = (s.semaine_fin - s.semaine_debut).dt.days + 1
    return s


def tests_calendrier(s, h, nm):
    res = []
    res.append(("chaque semaine dure 7 jours", bool((s.nb_jours == 7).all())))
    trous = (s.semaine_debut - s.semaine_fin.shift(1)).dt.days.dropna()
    res.append(("semaines consécutives (ni trou ni chevauchement)",
                bool((trous == 1).all())))
    res.append((f"même jour de début ({s.jour_debut.unique()[0]}) et de fin "
                f"({s.jour_fin.unique()[0]}) pour toutes les semaines",
                s.jour_debut.nunique() == 1 and s.jour_fin.nunique() == 1))
    # La normale écrite dans le bulletin = normale du jour de FIN de semaine
    cle = nm.set_index(["departement", "date"])["cumul_normale_mm"]
    hh = h.drop_duplicates(["semaine_debut", "departement"]).copy()
    hh["fin"] = pd.to_datetime(hh.semaine_fin).dt.date
    ok = [abs(cle[(r.departement, r.fin)] - r.normale_mm) <= TOL
          for r in hh.itertuples()]
    res.append((f"normale du bulletin = normale à la date de fin de semaine "
                f"({sum(ok)}/{len(ok)})", all(ok)))
    return res


# --------------------------------------------------------------------------
def main(chemin_xlsx, dossier):
    dossier = Path(dossier)
    moteur = db.obtenir_moteur(dossier)
    h = pd.read_sql("select * from observations_hebdo", moteur)
    wb = load_workbook(chemin_xlsx, data_only=True)

    a, moyennes = lire_annuel(wb["Pluvio 14 dernières années"])
    nm = lire_normales(wb["Normales cumulées"])
    s = construire_calendrier(h)

    tous = (tests_annuel(a, moyennes) + tests_normales(nm)
            + tests_calendrier(s, h, nm))

    # Enregistrement : dates au format texte ISO (AAAA-MM-JJ)
    nm_out = nm.assign(date=nm.date.astype(str))
    s_out = s.assign(semaine_debut=s.semaine_debut.dt.strftime("%Y-%m-%d"),
                     semaine_fin=s.semaine_fin.dt.strftime("%Y-%m-%d"))
    with moteur.begin() as con:
        a.to_sql("annuel", con, if_exists="replace", index=False)
        nm_out.to_sql("normales", con, if_exists="replace", index=False)
        s_out.to_sql("semaines", con, if_exists="replace", index=False)
    a.to_csv(dossier / "annuel_2012_2025.csv", index=False)
    nm_out.to_csv(dossier / "normales_1991_2020.csv", index=False)
    s_out.to_csv(dossier / "calendrier_semaines.csv", index=False)

    print(f"annuel : {len(a)} lignes ({a.poste.nunique()} postes x "
          f"{a.annee.nunique()} ans, {int(a.mm.isna().sum())} manquantes)")
    print(f"normales : {len(nm)} lignes | semaines : {len(s)}\n")
    for nom, ok in tous:
        print(f"  [{'OK' if ok else 'ECHEC'}] {nom}")
    print("\nDonnées manquantes signalées :")
    for x in alertes:
        print("  -", x["detail"])


if __name__ == "__main__":
    main(sys.argv[1], Path(__file__).parent / "donnees_propres")
