"""
Étape 3 - Indicateurs pluviométriques.

Lit la base créée par nettoyage.py et nettoyage_reference.py, calcule :
  1. l'écart à la normale (mm et %), par poste / département / région
  2. l'écart à 2025 à la même date
  3. l'intensité (mm par jour de pluie)
  4. le rang de 2026 parmi les années 2012-2025
  5. les cumuls possiblement sous-estimés (jours « np »)
puis vérifie les résultats contre le tableau « Récap pluvio » du fichier.

Utilisation :
    python indicateurs.py chemin/vers/le_fichier.xlsx
"""

import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

import db


DOSSIER = Path(__file__).parent / "donnees_propres"
SEUIL_PCT_MM = 10.0  # pas de pourcentage si la référence est < 10 mm
TOL = 0.1            # tolérance d'arrondi des tests (mm)


def charger():
    moteur = db.obtenir_moteur()
    noms = [
        "observations_hebdo",
        "semaines",
        "normales",
        "annuel",
        "mesures_journalieres"
    ]

    t = {
        n: pd.read_sql(f"select * from {n}", moteur)
        for n in noms
    }

    return t


def pct(ecart, ref):
    """Pourcentage de l'écart, vide si la référence est trop petite."""
    return np.where(
        ref >= SEUIL_PCT_MM,
        100 * ecart / ref,
        np.nan
    )


# --------------------------------------------------------------------------
# 1-3 + 5. Indicateurs par poste et par semaine
# --------------------------------------------------------------------------
def indicateurs_poste(t):
    h = t["observations_hebdo"].rename(columns={
        "normale_mm": "normale_bulletin_mm",
        "ecart_normale_mm": "ecart_bulletin_mm"
    })

    # Harmonisation des types de dates avant la jointure.
    # Dans observations_hebdo, semaine_debut est lu comme datetime.date.
    # Dans semaines, semaine_debut est lu comme chaîne de caractères.
    s = t["semaines"][["num_semaine", "semaine_debut"]].copy()

    h["semaine_debut"] = pd.to_datetime(
        h["semaine_debut"]
    )

    s["semaine_debut"] = pd.to_datetime(
        s["semaine_debut"]
    )

    h = h.merge(
        s,
        on="semaine_debut"
    )

    nm = t["normales"].rename(columns={
        "date": "semaine_fin",
        "cumul_normale_mm": "normale_mm"
    })

    # Même principe : on harmonise les dates utilisées
    # pour la jointure avec les normales.
    h["semaine_fin"] = pd.to_datetime(
        h["semaine_fin"]
    )

    nm["semaine_fin"] = pd.to_datetime(
        nm["semaine_fin"]
    )

    h = h.merge(
        nm[["departement", "semaine_fin", "normale_mm"]],
        on=["departement", "semaine_fin"],
        how="left"
    )

    h["ecart_normale_mm"] = (
        h.cumul_2026_mm - h.normale_mm
    )

    h["ecart_normale_pct"] = pct(
        h.ecart_normale_mm,
        h.normale_mm
    )

    h["ecart_2025_mm"] = (
        h.cumul_2026_mm - h.cumul_2025_mm
    )

    h["ecart_2025_pct"] = pct(
        h.ecart_2025_mm,
        h.cumul_2025_mm
    )

    h["intensite_mm_par_jour"] = np.where(
        h.cumul_2026_nj > 0,
        h.cumul_2026_mm / h.cumul_2026_nj,
        np.nan
    )

    # Jours « np » (non parvenus) déjà survenus à la fin de chaque semaine
    j = t["mesures_journalieres"]

    npj = j[
        j.statut == "non_parvenu"
    ][["poste", "date"]]

    # Harmonisation des dates pour comparer correctement
    # avec semaine_fin.
    npj = npj.copy()
    npj["date"] = pd.to_datetime(
        npj["date"]
    )

    h["np_jours_cumules"] = [
        int(
            (
                (npj.poste == r.poste)
                & (npj.date <= r.semaine_fin)
            ).sum()
        )
        for r in h.itertuples()
    ]

    h["donnees_incompletes"] = (
        h.np_jours_cumules > 0
    )

    cols = [
        "poste",
        "departement",
        "num_semaine",
        "semaine_debut",
        "semaine_fin",
        "cumul_2026_mm",
        "cumul_2026_nj",
        "normale_mm",
        "ecart_normale_mm",
        "ecart_normale_pct",
        "cumul_2025_mm",
        "ecart_2025_mm",
        "ecart_2025_pct",
        "intensite_mm_par_jour",
        "np_jours_cumules",
        "donnees_incompletes",
        "ecart_bulletin_mm"
    ]

    return h[cols]


def agreger(p):
    """Départements et région : moyenne simple des postes, comme le bulletin."""

    reg = p.assign(
        departement="Région"
    )

    tout = pd.concat([
        p.assign(niveau="Département"),
        reg.assign(niveau="Région")
    ])

    g = tout.groupby(
        ["niveau", "departement", "num_semaine"],
        as_index=False
    )

    ag = g.agg(
        semaine_debut=("semaine_debut", "first"),
        semaine_fin=("semaine_fin", "first"),
        nb_postes=("poste", "nunique"),
        cumul_2026_mm=("cumul_2026_mm", "mean"),
        cumul_2026_nj=("cumul_2026_nj", "mean"),
        normale_mm=("normale_mm", "mean"),
        cumul_2025_mm=("cumul_2025_mm", "mean"),
        nb_postes_incomplets=("donnees_incompletes", "sum")
    )

    ag = ag.rename(columns={
        "departement": "nom"
    })

    ag["ecart_normale_mm"] = (
        ag.cumul_2026_mm - ag.normale_mm
    )

    ag["ecart_normale_pct"] = pct(
        ag.ecart_normale_mm,
        ag.normale_mm
    )

    ag["ecart_2025_mm"] = (
        ag.cumul_2026_mm - ag.cumul_2025_mm
    )

    ag["ecart_2025_pct"] = pct(
        ag.ecart_2025_mm,
        ag.cumul_2025_mm
    )

    ag["intensite_mm_par_jour"] = (
        ag.cumul_2026_mm / ag.cumul_2026_nj
    )

    return ag


# --------------------------------------------------------------------------
# 4. Rang de 2026 parmi 2012-2025
# --------------------------------------------------------------------------
def historique(a, niveau, nom):
    d = a.dropna(
        subset=["mm"]
    )

    if niveau == "Poste":
        d = d[d.poste == nom]

    elif niveau == "Département":
        d = d[d.departement == nom]

    y = d.groupby("annee").agg(
        mm=("mm", "mean"),
        nj=("jours_pluie", "mean")
    )

    y["intensite"] = y.mm / y.nj

    return y


def classer(serie, x):
    """Rang de x parmi la série + x (1 = valeur la plus basse)."""

    return dict(
        n_annees=len(serie) + 1,
        rang_du_plus_bas=int(
            1 + (serie < x).sum()
        ),
        moyenne_2012_2025=serie.mean(),
        ecart_type=serie.std(ddof=1),
        z=(x - serie.mean()) / serie.std(ddof=1),
        minimum=serie.min(),
        maximum=serie.max()
    )


def rangs(p, a):
    fin = p[
        p.num_semaine == p.num_semaine.max()
    ]

    lignes = []

    entites = (
        [("Poste", n) for n in sorted(fin.poste)]
        + [
            ("Département", d)
            for d in ["Bambey", "Diourbel", "Mbacké"]
        ]
        + [
            ("Région", "Région")
        ]
    )

    for niveau, nom in entites:

        if niveau == "Poste":
            f = fin[
                fin.poste == nom
            ]

        elif niveau == "Département":
            f = fin[
                fin.departement == nom
            ]

        else:
            f = fin

        hist = historique(
            a,
            niveau,
            nom
        )

        mm = f.cumul_2026_mm.mean()
        nj = f.cumul_2026_nj.mean()

        for ind, col, x in [
            ("cumul_mm", "mm", mm),
            ("jours_pluie", "nj", nj),
            (
                "intensite_mm_par_jour",
                "intensite",
                mm / nj
            )
        ]:
            lignes.append(
                dict(
                    niveau=niveau,
                    nom=nom,
                    indicateur=ind,
                    valeur_2026=x,
                    **classer(hist[col], x)
                )
            )

    return pd.DataFrame(lignes)


# --------------------------------------------------------------------------
# Tests
# --------------------------------------------------------------------------
def tests(p, ag, chemin_xlsx):
    res = []

    e = (
        p.ecart_normale_mm
        - p.ecart_bulletin_mm
    ).abs()

    # Certaines lignes peuvent avoir une normale manquante.
    # Elles ne doivent pas faire échouer le test avec un maximum = NaN.
    e_valides = e.dropna()

    if len(e_valides) > 0:
        ecart_max = e_valides.max()
    else:
        ecart_max = 0.0

    res.append(
        (
            f"écart à la normale recalculé = écart du bulletin "
            f"({len(p)} lignes, écart max {ecart_max:.3f} mm)",
            ecart_max <= TOL
        )
    )

    ws = load_workbook(
        chemin_xlsx,
        data_only=True
    )["Récap pluvio"]

    attendu = {
        ("Département", "Bambey"): ws["G4"].value,
        ("Département", "Diourbel"): ws["G5"].value,
        ("Département", "Mbacké"): ws["G6"].value,
        ("Région", "Région"): ws["G7"].value
    }

    fin = (
        ag[
            ag.num_semaine == ag.num_semaine.max()
        ]
        .set_index(
            ["niveau", "nom"]
        )
        .cumul_2026_mm
    )

    d = max(
        abs(fin[k] - v)
        for k, v in attendu.items()
    )

    res.append(
        (
            f"moyennes départements et région = Récap "
            f"(écart max {d:.3f} mm)",
            d <= 0.01
        )
    )

    return res


def main(chemin_xlsx):
    t = charger()

    p = indicateurs_poste(t)

    ag = agreger(p)

    r = rangs(
        p,
        t["annuel"]
    )

    ok = tests(
        p,
        ag,
        chemin_xlsx
    )

    moteur = db.obtenir_moteur()

    with moteur.begin() as con:

        p.to_sql(
            "indicateurs_poste_semaine",
            con,
            if_exists="replace",
            index=False
        )

        ag.to_sql(
            "indicateurs_agreges_semaine",
            con,
            if_exists="replace",
            index=False
        )

        r.to_sql(
            "rang_2026",
            con,
            if_exists="replace",
            index=False
        )

    p.to_csv(
        DOSSIER / "indicateurs_poste_semaine.csv",
        index=False
    )

    ag.to_csv(
        DOSSIER / "indicateurs_agreges_semaine.csv",
        index=False
    )

    r.to_csv(
        DOSSIER / "rang_2026.csv",
        index=False
    )

    print("TESTS")

    for nom, v in ok:
        print(
            f"  [{'OK' if v else 'ECHEC'}] {nom}"
        )


if __name__ == "__main__":
    main(sys.argv[1])